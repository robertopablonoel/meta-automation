"""Publish stage: build the weekly review workbook (.xlsx) from week.json and
upload it to Google Drive as a converted native Google Sheet.

Design (Cole, 2026-07-31 v2): the Review Board is a MEETING INSTRUMENT, not an
export. Human-first rows — a name, a plain-English "why it's here" sentence,
five money numbers (spend/purch/ROAS/CPA/CTR), fast links, verdict. Deep
metrics live far right, past the verdict, for when someone wants to dig.

Conversion preserves tabs, dropdown validation, conditional formats, frozen
panes and HYPERLINK formulas (verified live 2026-07-31). Deterministic build.

Drive layout (config drive.weekly_reviews_folder_id = 'Weekly Ad Reviews'):
  Weekly Ad Reviews / <week folder> / {Weekly Ad Review sheet, videos/, archive/}
Idempotent: re-runs UPDATE the sheet in place (link never changes) and skip
already-archived files. Auth via gauth.py (Cole's own OAuth client).
"""
from __future__ import annotations

import json
import os

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .common import week_dir

# Route = the deterministic proposed next action (human can override in-sheet).
# Ordered to match the config "routes" values so the dropdown mirrors buckets.
ROUTES = ["ABO-scale 10-20%", "ABO-throttle ~5%", "ABO-push ~5%",
          "Graveyard", "Keep-in-arena"]
ROUTE_COLORS = {
    "ABO-scale 10-20%": "C6EFCE", "ABO-throttle ~5%": "D9EAD3",
    "ABO-push ~5%": "BDD7EE", "Graveyard": "F8CBAD",
    "Keep-in-arena": "EFEFEF",
}

# Evolve bucket sections (Review Board order). Winners first, losers last.
BUCKET_SECTIONS = [
    ("breakthrough", "🚀 BREAKTHROUGHS", "1E7A46",
     "funded AND scaling — push budget"),
    ("spend_winner", "💰 SPEND WINNERS", "2E7D32",
     "funded AND profitable but below scale KPI — throttle, protect margin"),
    ("kpi_winner", "🎯 KPI WINNERS", "1F6FB2",
     "not funded yet but hitting scaling ROAS — push to prove out"),
    ("loser", "⚰️ LOSERS → GRAVEYARD", "B23A3A",
     "funded-but-bleeding or spent-without-return — pull; mine any signal badge first"),
]

# palette
DARK = "212B3B"
ACCENT_SCALE = "1E7A46"
ACCENT_CREATIVE = "B07C10"
BAND = "F4F6FA"
ROAS_GREEN = "C6EFCE"
ROAS_AMBER = "FFF2CC"
ROAS_RED = "F8CBAD"

H = Font(bold=True, size=10, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor=DARK)
TITLE_F = Font(bold=True, size=16, color=DARK)
SUB_F = Font(bold=True, size=11)
MUTED_F = Font(size=9, color="808080")
WRAP = Alignment(wrap_text=True, vertical="top")
VC = Alignment(vertical="center")
WRAP_C = Alignment(wrap_text=True, vertical="center")
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))

PCT = "0.0%"
MONEY = "$#,##0"
NUM2 = "0.00"

# rollup ROAS traffic-light threshold; set from config.breakeven_roas in run().
_ROLLUP_BE = 1.40


def human(s: str | None) -> str:
    """'provocative_potency_warning' -> 'Provocative potency warning'."""
    if not s:
        return ""
    t = s.replace("_", " ").strip()
    return t[:1].upper() + t[1:]


def _ctr(v):  # link_ctr stored as percent number (3.88) -> fraction for PCT fmt
    return None if v is None else v / 100.0


def _hl(url, label):
    if not url:
        return ""
    return f'=HYPERLINK("{url}","{(label or url).replace(chr(34), chr(39))}")'


# ------------------------------------------------------------------ voice

def _fmt_pct(v):
    return f"{v * 100:.0f}%"


def why_text(x: dict, be: float) -> str:
    """One plain-English sentence: what earned this row its place + what to do."""
    m = x["metrics"]
    p, td = set(x["paths"]), set(x["top_decile"])
    bucket = x.get("bucket")
    bits = []
    if bucket == "breakthrough":
        bits.append(f"Funded & scaling — ${m['spend']:,.0f} in, "
                    f"{m['purchases']:.0f} purchases at {m['roas']:.2f}x (Meta feeding it hard)")
    elif bucket == "spend_winner":
        bits.append(f"Funded & profitable — ${m['spend']:,.0f} in at {m['roas']:.2f}x, "
                    f"below scale KPI (throttle)")
    elif bucket == "kpi_winner":
        bits.append(f"Hitting scale ROAS ({m['roas']:.2f}x) on ${m['spend']:,.0f} — "
                    f"not funded yet, push it")
    elif bucket == "loser" and x.get("bleeder"):
        roas_s = f"{m['roas']:.2f}x" if m["roas"] is not None else "no return"
        bits.append(f"BLEEDER — funded (${m['spend']:,.0f}) but {roas_s} (pull it)")
    if "hook" in p or "hook" in td:
        tag = " (top 10% this week)" if "hook" in td else " (top quartile)"
        bits.append(f"stops {_fmt_pct(m['thumbstop'])} of scrollers{tag}")
    if "hold" in p or "hold" in td:
        tag = " (top 10%)" if "hold" in td else ""
        bits.append(f"keeps {_fmt_pct(m['hold_rate'])} watching past the midpoint{tag}")
    if "ctr" in p or "ctr" in td:
        tag = " (top 10%)" if "ctr" in td else ""
        bits.append(f"{m['link_ctr']:.1f}% click through{tag}")
    if "efficiency" in p or "efficiency" in td:
        bits.append(f"{m['roas']:.1f}x ROAS on just ${m['spend']:,.0f}")
    if "intent" in p or "intent" in td:
        bits.append(f"carts piling up ({m['atc']:.0f} add-to-carts on ${m['spend']:,.0f})")
    if "algorithm_favorite" in p:
        bits.append("Meta keeps choosing to feed it budget")
    if bucket == "loser" and not x.get("bleeder") and m["roas"] is not None and m["roas"] < be:
        bits.append(f"not profitable ({m['roas']:.2f}x)")
    sentence = "; ".join(bits[:3]) if bits else "on the board by combined signal"
    action = (x.get("diagnosis") or {}).get("text")
    if action:
        sentence += f" → {action[:1].lower()}{action[1:]}"
    if (x.get("diagnosis") or {}).get("trending"):
        sentence += " ⚡hold rate rising — leads CPA by ~5-7 days"
    return sentence


def who_what(x: dict) -> str:
    p = x["name_parsed"]
    label = p.get("label") or ""
    if x["is_video"]:
        fmt = "🎬"
    else:
        fmt = "🖼"
    if label and label not in ("STAT", "VID"):
        return f"{fmt} 🤝 {label}"  # creator partnership ad
    return f"{fmt} brand"


def ad_display_name(x: dict) -> str:
    p = x["name_parsed"]
    concept = human(p.get("concept"))
    hook = human(p.get("hook"))
    if concept and hook:
        return f"{concept} — “{hook}…”"
    return x["ad_name"]


# ------------------------------------------------------------------ tabs

def build_brief(ws, w, cfg):
    t, c = w["totals"], w["counts"]
    win = w["window"]
    ws.column_dimensions["A"].width = 110
    ws.sheet_view.showGridLines = False
    rows = [
        (f"SPICY CUBES — WEEKLY AD REVIEW  {w['week']}  ({win['since']} → {win['until']})", TITLE_F),
        ("", None),
        (f"${t['spend']:,.0f} spend · {t['purchases']:.0f} purchases · ${t['revenue']:,.0f} revenue "
         f"· account-attributed ROAS {t['roas']:.2f} (breakeven {cfg['breakeven_roas']:.2f}, "
         f"scale {cfg['scaling_roas']:.2f})", SUB_F),
        (f"{t['ads_with_spend']} ads with spend → {t['unique_creatives']} unique creatives → "
         f"{c['review_rows']} on the Review Board  "
         f"(🚀 {c['breakthrough']} breakthrough · 💰 {c['spend_winner']} spend-winner · "
         f"🎯 {c['kpi_winner']} kpi-winner · ⚰️ {c['loser']} loser)", None),
        ("", None),
        ("THE MEETING IN ONE LINE: work the Review Board top to bottom, watch anything you don't "
         "recognize, set a verdict on every row. Verdicts become next week's briefs.", SUB_F),
    ]
    breakthroughs = [x for x in w["creatives"] if x["bucket"] == "breakthrough"]
    board = [x for x in w["creatives"] if x["bucket"]]
    rows.append(("", None))
    rows.append(("HIGHLIGHTS", SUB_F))
    if breakthroughs:
        top = breakthroughs[0]
        rows.append((f"• Biggest breakthrough: {ad_display_name(top)} — ${top['metrics']['spend']:,.0f}, "
                     f"{top['metrics']['purchases']:.0f} purchases, {top['metrics']['roas']:.2f}x", None))
    hookers = sorted([x for x in board if "hook" in x["paths"]],
                     key=lambda x: -(x["metrics"]["thumbstop"] or 0))[:2]
    for x in hookers:
        rows.append((f"• Hook to steal: {ad_display_name(x)} — {why_text(x, cfg['breakeven_roas'])}", None))
    conc = [r for r in w["rollups"]["concept"] if r["spend"] >= 300]
    if conc:
        best = max(conc, key=lambda r: r["roas"])
        worst = min(conc, key=lambda r: r["roas"])
        rows.append((f"• Best concept ≥$300: {human(best['label'])} — {best['roas']:.2f}x on ${best['spend']:,.0f}", None))
        if worst is not best:
            rows.append((f"• Weakest concept ≥$300: {human(worst['label'])} — {worst['roas']:.2f}x on ${worst['spend']:,.0f}", None))
    for i, (txt, f) in enumerate(rows, 1):
        cell = ws.cell(row=i, column=1, value=txt)
        cell.alignment = WRAP
        if f:
            cell.font = f


REVIEW_COLS = [
    ("#", 4), ("Ad", 38), ("Who", 15), ("Why it's on the board", 62),
    ("Spend", 9), ("$/day", 8), ("Purch", 7), ("ROAS", 7), ("CPA", 8),
    ("CTR", 7), ("Click %", 8),
    ("Watch", 9), ("Metrics", 9),
    # two human columns (was one VERDICT): proposed Route (dropdown, override) +
    # avatar-first Learning free-text for Alex.
    ("Route", 17), ("Learning (avatar-first)", 40), ("Notes", 30), ("Owner", 9),
    # deep detail, past the human zone
    ("Thumbstop", 10), ("Hold", 8), ("Avg watch s", 9), ("Freq", 7),
    ("Lifetime spend", 11), ("Lifetime ROAS", 9), ("Batch", 7), ("Pool", 11),
    ("Campaigns", 18), ("Raw ad name", 46),
]
ROUTE_COL = 14   # Route (proposed, dropdown)
ROAS_COL = 8
N_COLS = len(REVIEW_COLS)


def _section_row(ws, r, text, color):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=N_COLS)
    cell = ws.cell(row=r, column=1, value=text)
    cell.font = Font(bold=True, size=12, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=color)
    cell.alignment = VC
    ws.row_dimensions[r].height = 24


def _window_days(w) -> int:
    from datetime import date
    since = date.fromisoformat(w["window"]["since"])
    until = date.fromisoformat(w["window"]["until"])
    return (until - since).days + 1


# number-format column groups (1-indexed against REVIEW_COLS above)
_MONEY_COLS = (5, 6, 9, 22)
_NUM2_COLS = (8, 21, 23)
_PCT_COLS = (10, 11, 18, 19)
_WRAP_COLS = (2, 4, 15, 16, 27)


def _click_share(x: dict):
    """Share of purchases that are 7d-click vs total. The raw pull does not
    break purchases into click-only vs view (attribution stays 7d_click+1d_view
    per Cole/Roberto — we do NOT change it). TODO: add a second insights pull
    with action_attribution_windows=['7d_click'] to populate this; until then
    the column exists for visibility but reads n/a."""
    return None  # TODO: populate click-only purchases; see docstring


def build_review(ws, w, cfg, human_input=None):
    be = cfg["breakeven_roas"]
    days = _window_days(w)
    human_input = human_input or {}
    ws.sheet_view.showGridLines = False
    for i, (c, wd) in enumerate(REVIEW_COLS, 1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font = H
        cell.fill = HDR_FILL
        cell.alignment = WRAP_C
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.freeze_panes = "C2"

    acct = (w.get("config", {}).get("ad_account_id")
            or cfg.get("ad_account_id", "")).replace("act_", "")

    r = 2
    data_rows = []

    def emit(rows, accent):
        nonlocal r
        band = False
        for i, x in enumerate(rows, 1):
            m, lt = x["metrics"], x.get("lifetime") or {}
            inst = max(x["instances"], key=lambda z: z["metrics"]["spend"] or 0)
            watch = x.get("instagram_permalink_url") or x.get("preview_link")
            metrics_url = (f"https://adsmanager.facebook.com/adsmanager/manage/ads"
                           f"?act={acct}&selected_ad_ids={inst['ad_id']}") if acct else None
            # human columns: Route pre-filled with the proposed route unless a
            # prior sheet already carried a human override; Learning/Notes/Owner
            # preserved from prior sheet if present.
            prior = human_input.get(x["ad_name"], {})
            route = prior.get("route") or x.get("route") or ""
            click = _click_share(x)
            click_val = click if click is not None else "n/a"
            why = why_text(x, be)
            if x.get("batch_loser_pick"):
                why = "🔎 REVIEW-THIS-LOSER (batch pick) — " + why
            vals = [
                i, ad_display_name(x), who_what(x), why,
                m["spend"], (m["spend"] or 0) / days,
                m["purchases"], m["roas"], m["cpa"], _ctr(m["link_ctr"]), click_val,
                _hl(watch, "▶ watch"), _hl(metrics_url, "📊 open"),
                route, prior.get("learning", ""), prior.get("notes", ""), prior.get("owner", ""),
                (m["thumbstop"] if x["is_video"] else None),
                (m["hold_rate"] if x["is_video"] else None),
                (m["avg_watch_seconds"] if x["is_video"] else None),
                m["frequency"], lt.get("spend"), lt.get("roas"),
                "/".join(str(b) for b in x["batches"]) if x["batches"] else "",
                x["pool"],
                ", ".join(sorted({c.split(" - ")[0].replace("Campaign ", "") for c in x["campaigns"]})),
                x["ad_name"],
            ]
            for ci, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=ci, value=v)
                cell.border = THIN
                cell.alignment = WRAP if ci in _WRAP_COLS else VC
                if ci in _MONEY_COLS:
                    cell.number_format = MONEY
                if ci in _NUM2_COLS:
                    cell.number_format = NUM2
                if ci in _PCT_COLS:
                    cell.number_format = PCT
                if band and ci <= N_COLS:
                    if cell.fill.fgColor.rgb in (None, "00000000"):
                        cell.fill = PatternFill("solid", fgColor=BAND)
                if ci == 1:
                    cell.font = Font(bold=True, size=9, color=accent)
                # ROAS traffic light
                if ci == ROAS_COL and v is not None:
                    color = ROAS_GREEN if v > be else (ROAS_AMBER if v >= 1.0 else ROAS_RED)
                    cell.fill = PatternFill("solid", fgColor=color)
            ws.row_dimensions[r].height = 42
            data_rows.append(r)
            band = not band
            r += 1

    by_bucket = {b: [x for x in w["creatives"] if x["bucket"] == b]
                 for b, *_ in BUCKET_SECTIONS}
    for bucket, title, accent, blurb in BUCKET_SECTIONS:
        rows = by_bucket[bucket]
        if bucket == "loser":
            # only surface losers with meaningful spend; mark the per-batch
            # highest-spend loser as the "review this loser" pick.
            rows = [x for x in rows if x.get("batch_loser_pick")
                    or (x["metrics"]["spend"] or 0) >= cfg["funded"]["min_spend_abs"]]
        _section_row(ws, r, f"  {title} ({len(rows)}) — {blurb}", accent)
        r += 1
        # tag the batch review-pick loser inline in the Why column via row emit;
        # emit handles values, so pre-annotate the display name for picks.
        emit(rows, accent)
        r += 1  # spacer

    if data_rows:
        first, last = min(data_rows), max(data_rows)
        rcol = get_column_letter(ROUTE_COL)
        dv = DataValidation(type="list", formula1='"' + ",".join(ROUTES) + '"',
                            allow_blank=True, showDropDown=False)
        ws.add_data_validation(dv)
        dv.add(f"{rcol}{first}:{rcol}{last}")
        for route, color in ROUTE_COLORS.items():
            ws.conditional_formatting.add(
                f"{rcol}{first}:{rcol}{last}",
                FormulaRule(formula=[f'${rcol}{first}="{route}"'],
                            fill=PatternFill("solid", fgColor=color)))


def _rollup_block(ws, start, title, rows, label_hdr, days, humanize=True):
    ws.cell(row=start, column=1, value=title).font = Font(bold=True, size=12, color=DARK)
    winners = bool(rows) and "n_breakthrough" in rows[0]
    hdr = [label_hdr, "Ads", "Spend", "$/day", "Purch", "Revenue", "ROAS", "CPA"]
    if winners:
        hdr += ["🚀", "💰", "🎯", "⚰️"]
    for i, c in enumerate(hdr, 1):
        cell = ws.cell(row=start + 1, column=i, value=c)
        cell.font = H
        cell.fill = HDR_FILL
    r = start + 2
    band = False
    for x in rows:
        label = human(x["label"]) if humanize else (x["label"] or "")
        vals = [label, x.get("n_ads", x.get("n_creatives")), x["spend"],
                (x["spend"] or 0) / days,
                x["purchases"], x["revenue"], x["roas"], x["cpa"]]
        if winners:
            vals += [x["n_breakthrough"] or "", x["n_spend_winner"] or "",
                     x["n_kpi_winner"] or "", x["n_loser"] or ""]
        for i, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=i, value=v)
            cell.border = THIN
            if band:
                cell.fill = PatternFill("solid", fgColor=BAND)
            if i in (3, 4, 6, 8):
                cell.number_format = MONEY
            if i == 7:
                cell.number_format = NUM2
                if v is not None:
                    color = ROAS_GREEN if v > _ROLLUP_BE else (ROAS_AMBER if v >= 1.0 else ROAS_RED)
                    cell.fill = PatternFill("solid", fgColor=color)
        band = not band
        r += 1
    return r + 2


def build_rollups(ws, w, keys_titles):
    ws.sheet_view.showGridLines = False
    days = _window_days(w)
    ws.column_dimensions["A"].width = 52
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 10
    for col in "IJKL":
        ws.column_dimensions[col].width = 5
    r = 1
    for key, label in keys_titles:
        rows = sorted(w["rollups"][key], key=lambda x: -(x["spend"] or 0))
        r = _rollup_block(ws, r, f"PERFORMANCE BY {label.upper()}", rows, label, days,
                          humanize=(key in ("concept", "hook")))


ALL_COLS = [("Creative", 60), ("Bucket", 13), ("Route (proposed)", 17),
            ("Pool", 11), ("Ads", 5), ("Spend", 10), ("Spend %", 8),
            ("Purch", 7), ("Revenue", 10), ("ROAS", 7),
            ("CPA", 9), ("CTR", 8), ("Thumbstop", 9), ("Hold", 8),
            ("Signal badges", 22), ("Lifetime spend", 11), ("Lifetime ROAS", 9)]


def build_all_ads(ws, w):
    for i, (c, wd) in enumerate(ALL_COLS, 1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font = H
        cell.fill = HDR_FILL
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.freeze_panes = "A2"
    r = 2
    for x in w["creatives"]:
        m, lt = x["metrics"], x.get("lifetime") or {}
        bucket = (x.get("bucket") or "").upper()
        if x.get("bleeder"):
            bucket += " (BLEEDER)"
        vals = [x["ad_name"], bucket, x.get("route") or "", x["pool"], x["n_instances"],
                m["spend"], x.get("spend_share"), m["purchases"], m["revenue"],
                m["roas"], m["cpa"], _ctr(m["link_ctr"]),
                (m["thumbstop"] if x["is_video"] else None),
                (m["hold_rate"] if x["is_video"] else None),
                ", ".join(x.get("signal_badges", x.get("paths", []))),
                lt.get("spend"), lt.get("roas")]
        for i, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=i, value=v)
            if i in (6, 9, 16):          # Spend, Revenue, Lifetime spend
                cell.number_format = MONEY
            if i in (10, 17):            # ROAS, Lifetime ROAS
                cell.number_format = NUM2
            if i in (7, 12, 13, 14):     # Spend %, CTR, Thumbstop, Hold
                cell.number_format = PCT
        r += 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(ALL_COLS))}{r-1}"


def build_verdict_log(ws, week):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 110
    ws.cell(row=1, column=1, value="LAST WEEK'S VERDICTS — follow-through").font = TITLE_F
    cell = ws.cell(row=3, column=1,
                   value=f"{week} is the FIRST weekly review — no prior verdicts to show. "
                         "From next week this tab lists every verdict set last Friday, "
                         "who owned it, and what actually happened "
                         "(executed / pending / superseded).")
    cell.alignment = WRAP


def fetch_existing_verdicts(cfg: dict, week: str, w: dict) -> dict:
    """Read the human columns (Route, Learning, Notes, Owner) from an already-
    published week sheet, keyed by raw ad name, so a republish never wipes human
    input. Backward-tolerant: a prior sheet with the old single VERDICT column
    (no Route/Learning) still restores what it has (VERDICT -> Notes prefix).

    Returns {raw_ad_name: {route, learning, notes, owner}}."""
    drive_cfg = cfg.get("drive") or {}
    root = drive_cfg.get("weekly_reviews_folder_id")
    if not root:
        return {}
    try:
        s = _drive_session()
        wf = _find(s, _week_folder_name(week, w), root, FOLDER_MIME)
        if not wf:
            return {}
        sid = _find(s, f"Weekly Ad Review — {week}", wf, GSHEET_MIME)
        if not sid:
            return {}
        r = s.get(f"https://sheets.googleapis.com/v4/spreadsheets/{sid}/values/"
                  f"'Review Board'!A1:AZ2000", timeout=60)
        r.raise_for_status()
        rows = r.json().get("values", [])
    except Exception as e:
        print(f"[publish] WARN could not read existing human columns ({e})")
        return {}
    if not rows:
        return {}
    hdr = rows[0]

    def col(name):
        return hdr.index(name) if name in hdr else None

    ci_route = col("Route")
    ci_learn = col("Learning (avatar-first)")
    ci_notes = col("Notes")
    ci_owner = col("Owner")
    ci_verdict = col("VERDICT")            # legacy sheets only
    ci_raw = col("Raw ad name")
    if ci_raw is None:
        return {}
    out = {}
    for row in rows[1:]:
        def get(ci):
            return (row[ci].strip() if ci is not None and ci < len(row) and row[ci] else "")
        raw = get(ci_raw)
        if not raw:
            continue
        route = get(ci_route)
        learning = get(ci_learn)
        notes = get(ci_notes)
        owner = get(ci_owner)
        legacy = get(ci_verdict)
        if legacy and not route and not learning:
            # migrate an old VERDICT value into Notes so it isn't lost
            notes = (f"[prior verdict: {legacy}] " + notes).strip()
        if route or learning or notes or owner:
            out[raw] = {"route": route, "learning": learning,
                        "notes": notes, "owner": owner}
    if out:
        print(f"[publish] preserving {len(out)} human-edited rows "
              "(Route/Learning/Notes/Owner)")
    return out


def run(cfg: dict, week: str) -> None:
    global _ROLLUP_BE
    _ROLLUP_BE = cfg["breakeven_roas"]
    d = week_dir(cfg, week)
    with open(os.path.join(d, "week.json")) as f:
        w = json.load(f)

    human_input = fetch_existing_verdicts(cfg, week, w)

    wb = Workbook()
    ws = wb.active
    if cfg.get("include_brief", True):
        ws.title = "Meeting Brief"
        build_brief(ws, w, cfg)
        rb = wb.create_sheet("Review Board")
    else:
        ws.title = "Review Board"
        rb = ws
    build_review(rb, w, cfg, human_input)
    # CONCEPT is the primary rollup read now that an ad set ≈ one concept:
    # "which concept won" leads its own tab; batch/hook follow.
    build_rollups(wb.create_sheet("Concepts (primary)"),
                  w, [("concept", "Concept")])
    build_rollups(wb.create_sheet("Batches & Hooks"),
                  w, [("batch", "Batch"), ("hook", "Hook")])
    build_rollups(wb.create_sheet("Ad Sets & Campaigns"),
                  w, [("campaign", "Campaign"), ("adset", "Ad Set")])
    build_all_ads(wb.create_sheet("All Ads"), w)
    # 'Last Week's Verdicts' tab only exists once there ARE prior verdicts to
    # show (W32+; read back via fetch_existing_verdicts on the prior week).
    # Cole deleted the empty placeholder from W31 — don't resurrect it.

    out = os.path.join(d, f"review_{week}.xlsx")
    wb.save(out)
    print(f"[publish] wrote {out}")
    upload_to_drive(cfg, week, w, out)


# ------------------------------------------------------------------ Drive

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
GSHEET_MIME = "application/vnd.google-apps.spreadsheet"
FOLDER_MIME = "application/vnd.google-apps.folder"


def _drive_session():
    import requests as rq

    from .gauth import get_credentials
    creds = get_credentials()
    s = rq.Session()
    s.headers["Authorization"] = f"Bearer {creds.token}"
    return s


def _find(s, name, parent, mime=None):
    q = f"name = '{name}' and '{parent}' in parents and trashed = false"
    if mime:
        q += f" and mimeType = '{mime}'"
    r = s.get("https://www.googleapis.com/drive/v3/files",
              params={"q": q, "fields": "files(id,name)"}, timeout=60)
    r.raise_for_status()
    files = r.json().get("files", [])
    return files[0]["id"] if files else None


def _ensure_folder(s, name, parent):
    fid = _find(s, name, parent, FOLDER_MIME)
    if fid:
        return fid
    r = s.post("https://www.googleapis.com/drive/v3/files",
               json={"name": name, "mimeType": FOLDER_MIME, "parents": [parent]},
               timeout=60)
    r.raise_for_status()
    return r.json()["id"]


def _upload(s, path, name, parent, src_mime, dest_mime=None, existing_id=None):
    meta = {"name": name}
    if dest_mime:
        meta["mimeType"] = dest_mime
    if existing_id:
        url = f"https://www.googleapis.com/upload/drive/v3/files/{existing_id}?uploadType=multipart"
        method = s.patch
    else:
        meta["parents"] = [parent]
        url = "https://www.googleapis.com/upload/drive/v3/files?uploadType=multipart"
        method = s.post
    with open(path, "rb") as f:
        r = method(url, files={
            "metadata": ("metadata", json.dumps(meta), "application/json; charset=UTF-8"),
            "file": (name, f, src_mime),
        }, timeout=300)
    r.raise_for_status()
    return r.json()["id"]


def _week_folder_name(week: str, w: dict) -> str:
    from datetime import date
    since = date.fromisoformat(w["window"]["since"])
    until = date.fromisoformat(w["window"]["until"])
    if since.month == until.month:
        span = f"{since.strftime('%b')} {since.day}–{until.day}"
    else:
        span = f"{since.strftime('%b')} {since.day}–{until.strftime('%b')} {until.day}"
    return f"{week} ({span})"


def upload_to_drive(cfg: dict, week: str, w: dict, xlsx_path: str) -> None:
    drive_cfg = cfg.get("drive") or {}
    root = drive_cfg.get("weekly_reviews_folder_id")
    if not root:
        print("[publish] no drive.weekly_reviews_folder_id in config; skipping upload")
        return
    try:
        s = _drive_session()
    except Exception as e:
        print(f"[publish] WARN no Google token ({e}); sheet NOT uploaded")
        return

    wf = _ensure_folder(s, _week_folder_name(week, w), root)
    _ensure_folder(s, "videos", wf)
    archive = _ensure_folder(s, "archive", wf)

    sheet_name = f"Weekly Ad Review — {week}"
    existing = _find(s, sheet_name, wf, GSHEET_MIME)
    sid = _upload(s, xlsx_path, sheet_name, wf, XLSX_MIME, GSHEET_MIME, existing)
    action = "updated" if existing else "created"
    print(f"[publish] sheet {action}: https://docs.google.com/spreadsheets/d/{sid}")

    d = os.path.dirname(xlsx_path)
    for fn in sorted(os.listdir(d)):
        if not (fn.endswith(".json") or fn.endswith(".txt")):
            continue
        if _find(s, fn, archive):
            continue
        src_mime = "application/json" if fn.endswith(".json") else "text/plain"
        _upload(s, os.path.join(d, fn), fn, archive, src_mime)
        print(f"[publish] archived {fn}")
    print(f"[publish] week folder: https://drive.google.com/drive/folders/{wf}")
